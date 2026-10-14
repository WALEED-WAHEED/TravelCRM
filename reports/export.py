"""Excel (.xlsx) export of the Reports page: one sheet per section."""
from io import BytesIO

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analytics

MONEY = '#,##0'
HEAD_FILL = PatternFill("solid", fgColor="1F5FBF")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def _sheet(wb, title, rtl):
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.rightToLeft = rtl
    return ws


def _header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=str(label))
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _rows(ws, start_row, rows, money_cols=()):
    r = start_row
    for r, values in enumerate(rows, start_row):
        for col, v in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=float(v) if hasattr(v, "as_tuple") else v)
            if col in money_cols:
                c.number_format = MONEY
    return r


def _autosize(ws, min_width=10, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))


def _change(delta):
    if not delta:
        return "—"
    sign = {"up": "+", "down": "−"}.get(delta["dir"], "")
    return f"{sign}{delta['pct']:.0f}%"


def build_xlsx(agency, period, report, rtl=False):
    wb = Workbook()
    wb.remove(wb.active)
    k, prev, deltas = report["kpis"], report["previous"], report["deltas"]
    fmt = "%d %b %Y"

    # ---- Summary ------------------------------------------------------------
    ws = _sheet(wb, _("Summary"), rtl)
    ws["A1"] = f"{agency.name} — {_('Report')}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{_('Period')}: {period['start']:{fmt}} – {period['end']:{fmt}}"
    ws["A3"] = f"{_('Compared with')}: {period['prev_start']:{fmt}} – {period['prev_end']:{fmt}}"
    from member import attendance as att_rules
    ws["A4"] = f"{_('Generated')}: {att_rules.office_now(att_rules.get_settings(agency)):{fmt} %I:%M %p}"
    f = period.get("filters") or {}
    if f:
        parts = []
        if f.get("staff"):
            from accounts.models import User
            u = User.objects.filter(pk=f["staff"]).first()
            parts.append(f"{_('Staff')}: {u.get_full_name() or u.username if u else f['staff']}")
        if f.get("package"):
            parts.append(f"{_('Package')}: {f['package']}")
        if f.get("party"):
            parts.append(f"{_('Party')}: {f['party']}")
        ws["A5"] = f"{_('Filters')}: " + " · ".join(parts)
    _header(ws, 6, [_("Measure"), _("This period"), _("Previous period"), _("Change")])
    measures = [
        ("bookings", _("Bookings"), False), ("booking_value", _("Booking value"), True),
        ("avg_value", _("Average booking"), True), ("pax", _("PAX"), False),
        ("cancelled", _("Cancelled"), False), ("cash_in", _("Money in"), True), ("cash_out", _("Money out"), True),
        ("net_cash", _("Net cash flow"), True), ("expenses", _("Expenses"), True),
        ("sales", _("Sales charged"), True), ("purchases", _("Supplier bills"), True), ("margin", _("Gross margin"), True),
    ]
    rows = [[label, k[key], prev[key], _change(deltas.get(key))] for key, label, _m in measures]
    att = report["attendance"]["kpis"]
    rows.append([_("Attendance rate (%)"), att["rate"], None, _change(report["attendance"]["deltas"].get("rate"))])
    rows.append([_("Parties owe you"), report["receivable_total"], None, ""])
    rows.append([_("You owe parties"), report["payable_total"], None, ""])
    last = _rows(ws, 7, rows, money_cols=(2, 3))
    if report["insights"]:
        r = last + 2
        ws.cell(row=r, column=1, value=_("What stands out")).font = BOLD
        for i, (_tone, text) in enumerate(report["insights"], r + 1):
            ws.cell(row=i, column=1, value=f"• {text}")
    _autosize(ws)
    ws.column_dimensions["A"].width = 34

    # ---- Over time ------------------------------------------------------------
    ws = _sheet(wb, _("Over time"), rtl)
    a_series = report["series"]["attendance"]
    a = dict(zip(a_series["keys"], a_series["values"], strict=True))
    _header(ws, 1, [_("Period"), _("Bookings"), _("Cancelled"), _("Booking value"), _("PAX"), _("Money in"),
                    _("Money out"), _("Net cash"), _("Expenses"), _("On time"), _("Late"), _("Absent")])
    rows = []
    for i, row in enumerate(analytics.timeseries_rows(report)):
        rows.append(row + [a.get("present", [0])[i], a.get("late", [0])[i], a.get("absent", [0])[i]])
    _rows(ws, 2, rows, money_cols=(4, 6, 7, 8, 9))
    _autosize(ws)

    # ---- Bookings --------------------------------------------------------------
    ws = _sheet(wb, _("Bookings"), rtl)
    _header(ws, 1, [_("Date"), _("Voucher"), _("Client"), _("Package"), _("Staff"), _("PAX"), _("Status"), _("Value")])
    last = _rows(ws, 2, [[b["date"], b["voucher"], b["client"], b["package"], b["staff"], b["pax"], b["status_label"], b["value"]]
                         for b in report["booking_rows"]], money_cols=(8,))
    for r in range(2, last + 1):
        ws.cell(row=r, column=1).number_format = "dd mmm yyyy"
    _autosize(ws)

    # ---- Party balances (as of the period end) -----------------------------------
    ws = _sheet(wb, _("Party balances"), rtl)
    _header(ws, 1, [_("Party"), _("Debit"), _("Credit"), _("Balance"), _("Status"), _("Last entry")])
    rows = []
    for b in report["party_rows"]:
        status = _("Owes you") if b["balance"] > 0 else _("You owe") if b["balance"] < 0 else _("Settled")
        rows.append([b["name"], b["debit"], b["credit"], abs(b["balance"]), status, b["last_date"]])
    last = _rows(ws, 2, rows, money_cols=(2, 3, 4))
    for r in range(2, last + 1):
        ws.cell(row=r, column=6).number_format = "dd mmm yyyy"
    _autosize(ws)

    # ---- Breakdowns -------------------------------------------------------------
    ws = _sheet(wb, _("Breakdowns"), rtl)
    b = report["breakdowns"]
    r = 1

    def block(title, heads, rows, money_cols):
        nonlocal r
        ws.cell(row=r, column=1, value=title).font = BOLD
        for col, h in enumerate(heads, 1):
            c = ws.cell(row=r + 1, column=col, value=str(h))
            c.fill, c.font = HEAD_FILL, HEAD_FONT
        end = _rows(ws, r + 2, rows, money_cols) if rows else r + 1
        if not rows:
            ws.cell(row=r + 2, column=1, value=_("No data"))
            end = r + 2
        r = end + 2

    s = b["services"]
    block(_("Sales vs purchases by service"), [_("Service"), _("Sold"), _("Bought")],
          list(zip(s["labels"], s["sales"], s["purchases"], strict=True)), (2, 3))
    block(_("Expenses by category"), [_("Category"), _("Amount")],
          list(zip(b["expenses"]["labels"], b["expenses"]["values"], strict=True)), (2,))
    block(_("Money in by payment method"), [_("Method"), _("Amount")],
          list(zip(b["methods"]["labels"], b["methods"]["values"], strict=True)), (2,))
    st = b["staff"]
    block(_("Sales by staff"), [_("Staff"), _("Booking value"), _("Bookings")],
          list(zip(st["labels"], st["values"], st["counts"], strict=True)), (2,))
    block(_("Top clients"), [_("Client"), _("Booking value")],
          list(zip(b["clients"]["labels"], b["clients"]["values"], strict=True)), (2,))
    block(_("Who paid you the most"), [_("Party"), _("Amount")],
          list(zip(b["payers"]["labels"], b["payers"]["values"], strict=True)), (2,))
    block(_("Top packages"), [_("Package"), _("Bookings")],
          list(zip(b["packages"]["labels"], b["packages"]["values"], strict=True)), ())
    _autosize(ws)

    # ---- Attendance -------------------------------------------------------------
    ws = _sheet(wb, _("Attendance"), rtl)
    _header(ws, 1, [_("Staff"), _("Attendance rate (%)"), _("On time"), _("Late"), _("Absent"), _("Leave"), _("Hours")])
    _rows(ws, 2, [[p["name"], p["rate"], p["present"], p["late"], p["absent"], p["leave"], p["hours"]]
                  for p in report["attendance"]["people"]])
    _autosize(ws)

    # ---- Installments -------------------------------------------------------------
    ws = _sheet(wb, _("Installments"), rtl)
    _header(ws, 1, [_("Due"), _("Booking"), _("Client"), _("Installment"), _("Amount"), _("Status")])
    today = timezone.localdate()
    rows = []
    for i in report["installments"]:
        bk = i.payment_plan.booking
        rows.append([i.due_date, bk.voucher_number, bk.client_company, i.label, i.amount,
                     _("Overdue") if i.due_date < today else _("Due")])
    last = _rows(ws, 2, rows, money_cols=(5,))
    for r in range(2, last + 1):
        ws.cell(row=r, column=1).number_format = "dd mmm yyyy"
    _autosize(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
