"""Shared Excel (.xlsx) styling and sheet helpers for the insights exports."""

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = '#,##0'
HEAD_FILL = PatternFill("solid", fgColor="1F5FBF")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def _sheet(wb, title, rtl):
    ws = wb.create_sheet(title[:31])
    ws.sheet_view.rightToLeft = rtl
    return ws


def _header(ws, row, labels):
    for col, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=col, value=str(label))
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _rows(ws, start_row, rows, money_cols=()):
    r = start_row
    for r, values in enumerate(rows, start_row):
        for col, v in enumerate(values, 1):
            c = ws.cell(row=r, column=col, value=float(v) if hasattr(v, "as_tuple") else v)
            if col in money_cols:
                c.number_format = MONEY
    return r


def _autosize(ws, min_width=10, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_width, min(max_width, w + 2))




