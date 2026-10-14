"""Approval insights: how vouchers agents submit move through approval — volumes, outcomes, decision speed,
who submits and who decides, why vouchers are sent back — and the Excel export of it.

Range = vouchers submitted in the chosen period.
Outcomes are each voucher's current state: approved, sent back (waiting for its agent to fix it) or still pending.
"""
from collections import defaultdict

from django.db.models import Q
from django.utils import timezone
from django.utils.translation import gettext as _

from member import analytics

from .models import Booking, BookingHotel, BookingMutamer

OUTCOMES = ("approved", "rejected", "pending")
SAVED = ("confirmed", "completed")
LIST_LIMIT = 10


def _submitted(agency):
    """Vouchers that went through approval: waiting, sent back, or approved by someone."""
    return (Booking.objects.filter(agency=agency)
            .filter(Q(status__in=("pending", "rejected")) | Q(decided_by__isnull=False, status__in=SAVED))
            .select_related("agent", "decided_by"))


def _when(b):
    return b.submitted_at or b.created_at


def _in_range(qs, start, end):
    return qs.filter(
        Q(submitted_at__date__gte=start, submitted_at__date__lte=end)
        | Q(submitted_at__isnull=True, created_at__date__gte=start, created_at__date__lte=end))


def _outcome(b):
    return "pending" if b.status == "pending" else ("rejected" if b.status == "rejected" else "approved")


def _pax(b):
    return b.pax_adults + b.pax_children + b.pax_infants


def _name(user):
    return (user.get_full_name() or user.username) if user else "—"


def _hours(b):
    if b.decided_at and _when(b) and b.decided_at >= _when(b):
        return (b.decided_at - _when(b)).total_seconds() / 3600
    return None


def _duration(hours):
    if hours is None:
        return None
    if hours < 1:
        return _("%(n)s min") % {"n": max(round(hours * 60), 1)}
    if hours < 48:
        return _("%(n)s h") % {"n": round(hours, 1)}
    return _("%(n)s days") % {"n": round(hours / 24, 1)}


def _kpis(rows):
    total = len(rows)
    counts = {o: sum(1 for b in rows if _outcome(b) == o) for o in OUTCOMES}
    decided = counts["approved"] + counts["rejected"]
    hours = [h for h in (_hours(b) for b in rows if _outcome(b) != "pending") if h is not None]
    return {
        "submitted": total, "approved": counts["approved"], "rejected": counts["rejected"], "pending": counts["pending"],
        "approval_rate": round(counts["approved"] * 100 / decided) if decided else None,
        "avg_hours": round(sum(hours) / len(hours), 1) if hours else None,
        "avg_text": _duration(sum(hours) / len(hours)) if hours else None,
        "pax": sum(_pax(b) for b in rows),
    }


def _buckets(values, edges):
    return {"labels": [e[0] for e in edges],
            "values": [sum(1 for v in values if v >= lo and (hi is None or v < hi)) for _l, lo, hi in edges]}


def _pairs(counter, n=8):
    items = analytics._top(counter.items(), n)
    return {"labels": [str(k) for k, _v in items], "values": [float(v) for _k, v in items]}


def build(agency, period):
    start, end, g = period["start"], period["end"], period["granularity"]
    base = _submitted(agency)
    rows = list(_in_range(base, start, end).order_by("-submitted_at", "-created_at"))
    prev = list(_in_range(base, period["prev_start"], period["prev_end"]))
    k, pk = _kpis(rows), _kpis(prev)
    deltas = {key: analytics._delta(k[key] or 0, pk[key] or 0) if k[key] is not None else None
              for key in ("submitted", "approved", "rejected")}
    ids = [b.id for b in rows]

    # ---- over time -----------------------------------------------------------------------------------------------------
    starts = analytics.bucket_starts(start, end, g)
    labels = [analytics.bucket_label(d, g) for d in starts]
    index = {d: i for i, d in enumerate(starts)}
    by_outcome = {o: [0] * len(starts) for o in OUTCOMES}
    for b in rows:
        i = index.get(analytics.bucket_key(timezone.localtime(_when(b)).date(), g))
        if i is not None:
            by_outcome[_outcome(b)][i] += 1
    outcome_labels = {"approved": _("Approved"), "rejected": _("Sent back"), "pending": _("Pending")}

    # ---- decision speed ------------------------------------------------------------------------------------------------
    hours = [h for h in (_hours(b) for b in rows if _outcome(b) != "pending") if h is not None]
    speed = _buckets(hours, [(_("Under 1 hour"), 0, 1), (_("1–4 hours"), 1, 4), (_("4–24 hours"), 4, 24),
                             (_("1–3 days"), 24, 72), (_("Over 3 days"), 72, None)])

    # ---- people ---------------------------------------------------------------------------------------------------------
    agents, deciders = {}, {}
    for b in rows:
        a = agents.setdefault(_name(b.agent), {"submitted": 0, "approved": 0, "rejected": 0, "pending": 0, "hours": []})
        a["submitted"] += 1
        a[_outcome(b)] += 1
        h = _hours(b)
        if h is not None and _outcome(b) != "pending":
            a["hours"].append(h)
        if b.decided_by_id and _outcome(b) != "pending":
            d = deciders.setdefault(_name(b.decided_by), {"approved": 0, "rejected": 0, "hours": []})
            d[_outcome(b)] += 1
            if h is not None:
                d["hours"].append(h)

    def avg(vals):
        return round(sum(vals) / len(vals), 1) if vals else None

    agent_rows = sorted(({"agent": n, "submitted": a["submitted"], "approved": a["approved"], "rejected": a["rejected"],
                          "pending": a["pending"],
                          "rate": round(a["approved"] * 100 / (a["approved"] + a["rejected"])) if a["approved"] + a["rejected"] else None,
                          "avg_hours": avg(a["hours"])} for n, a in agents.items()),
                        key=lambda r: (-r["submitted"], r["agent"]))
    decider_rows = sorted(({"member": n, "approved": d["approved"], "rejected": d["rejected"],
                            "decisions": d["approved"] + d["rejected"], "avg_hours": avg(d["hours"])}
                           for n, d in deciders.items()), key=lambda r: (-r["decisions"], r["member"]))

    # ---- why vouchers come back ------------------------------------------------------------------------------------------
    reasons = [{"id": b.id, "voucher": b.voucher_number, "agent": _name(b.agent), "at": b.decided_at,
                "reason": b.rejection_reason, "head": ""}
               for b in sorted((b for b in rows if b.status == "rejected" and b.rejection_reason),
                               key=lambda b: b.decided_at or _when(b), reverse=True)]

    # ---- what agents submit ----------------------------------------------------------------------------------------------
    pk_rows = {}
    for b in rows:
        key = b.package_name.strip() if b.package_name else ""
        e = pk_rows.setdefault(key, {"submitted": 0, "approved": 0, "rejected": 0, "pending": 0, "pax": 0})
        e["submitted"] += 1
        e[_outcome(b)] += 1
        e["pax"] += _pax(b)
    package_rows = sorted(({"package": key or _("Custom trip"), "custom": not key, **e,
                            "rate": round(e["approved"] * 100 / (e["approved"] + e["rejected"])) if e["approved"] + e["rejected"] else None}
                           for key, e in pk_rows.items()), key=lambda r: (-r["submitted"], r["package"]))
    sizes = _buckets([_pax(b) for b in rows], [("1–2", 1, 3), ("3–5", 3, 6), ("6–10", 6, 11), ("11–20", 11, 21), ("21+", 21, None)])
    city_nights = defaultdict(int)
    for h in BookingHotel.objects.filter(booking_id__in=ids).values("city", "nights"):
        city_nights[(h["city"] or _("Other")).strip().title()] += h["nights"] or 0
    heads = {}
    for m in BookingMutamer.objects.filter(booking_id__in=[r["id"] for r in reasons]).order_by("id").values("booking_id", "name"):
        heads.setdefault(m["booking_id"], m["name"])
    for r in reasons:
        r["head"] = heads.get(r["id"], "")

    return {
        "kpis": k, "deltas": deltas,
        "tables": {"agents": agent_rows[:LIST_LIMIT], "deciders": decider_rows[:LIST_LIMIT],
                   "reasons": reasons[:LIST_LIMIT], "packages": package_rows[:LIST_LIMIT]},
        "full": {"agents": agent_rows, "deciders": decider_rows, "reasons": reasons, "packages": package_rows},
        "charts": {
            "labels": labels,
            "outcome": {"keys": list(OUTCOMES), "labels": [str(outcome_labels[o]) for o in OUTCOMES],
                        "values": [by_outcome[o] for o in OUTCOMES]},
            "speed": speed, "sizes": sizes, "cities": _pairs(city_nights),
        },
    }


def build_xlsx(agency, period, data, rtl=False):
    """Excel workbook with what the page shows: Summary, Trends, Agents, Decisions, Sent back, Packages, Group."""
    from io import BytesIO

    from openpyxl import Workbook

    from member.xlsx import BOLD, TITLE_FONT, _autosize, _header, _rows, _sheet

    k, c, t = data["kpis"], data["charts"], data["full"]
    wb = Workbook()
    wb.remove(wb.active)
    fmt = "%d %b %Y"

    def block(ws, row, title, headers, rows):
        ws.cell(row=row, column=1, value=title).font = BOLD
        _header(ws, row + 1, headers)
        end = _rows(ws, row + 2, rows) if rows else row + 1
        return max(end, row + 2) + 2

    def pairs(p):
        return list(zip(p["labels"], p["values"], strict=True))

    ws = _sheet(wb, _("Summary"), rtl)
    ws["A1"] = f"{agency.name} — {_('Approval insights')}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"{_('Period')}: {period['start']:{fmt}} – {period['end']:{fmt}} ({_('submitted date')})"
    ws["A3"] = f"{_('Compared with')}: {period['prev_start']:{fmt}} – {period['prev_end']:{fmt}}"
    _header(ws, 5, [_("Measure"), _("Value")])
    _rows(ws, 6, [[_("Submitted"), k["submitted"]], [_("Approved"), k["approved"]], [_("Sent back"), k["rejected"]],
                  [_("Still pending"), k["pending"]], [_("Approval rate (%)"), k["approval_rate"]],
                  [_("Average decision time (hours)"), k["avg_hours"]], [_("Travellers (PAX)"), k["pax"]]])
    _autosize(ws)
    ws.column_dimensions["A"].width = 34

    ws = _sheet(wb, _("Trends"), rtl)
    oc = dict(zip(c["outcome"]["keys"], c["outcome"]["values"], strict=True))
    row = block(ws, 1, _("Submitted by outcome"), [_("Period")] + c["outcome"]["labels"],
                [[label] + [oc[o][i] for o in c["outcome"]["keys"]] for i, label in enumerate(c["labels"])])
    block(ws, row, _("Decision time"), ["", _("Vouchers")], pairs(c["speed"]))
    _autosize(ws)

    ws = _sheet(wb, _("Agents"), rtl)
    _header(ws, 1, [_("Agent"), _("Submitted"), _("Approved"), _("Sent back"), _("Pending"), _("Approval rate (%)"), _("Average decision time (hours)")])
    _rows(ws, 2, [[r["agent"], r["submitted"], r["approved"], r["rejected"], r["pending"], r["rate"], r["avg_hours"]] for r in t["agents"]])
    _autosize(ws)

    ws = _sheet(wb, _("Decisions"), rtl)
    _header(ws, 1, [_("Team member"), _("Decisions"), _("Approved"), _("Sent back"), _("Average decision time (hours)")])
    _rows(ws, 2, [[r["member"], r["decisions"], r["approved"], r["rejected"], r["avg_hours"]] for r in t["deciders"]])
    _autosize(ws)

    ws = _sheet(wb, _("Sent back"), rtl)
    _header(ws, 1, [_("Voucher #"), _("Family head"), _("Agent"), _("Sent back on"), _("Reason")])
    last = _rows(ws, 2, [[r["voucher"], r["head"], r["agent"], r["at"].date() if r["at"] else None, r["reason"]] for r in t["reasons"]])
    for r in range(2, last + 1):
        ws.cell(row=r, column=4).number_format = "dd mmm yyyy"
    _autosize(ws)

    ws = _sheet(wb, _("Packages"), rtl)
    _header(ws, 1, [_("Package"), _("Submitted"), _("Approved"), _("Sent back"), _("Pending"), _("PAX"), _("Approval rate (%)")])
    _rows(ws, 2, [[r["package"], r["submitted"], r["approved"], r["rejected"], r["pending"], r["pax"], r["rate"]] for r in t["packages"]])
    _autosize(ws)

    ws = _sheet(wb, _("Group"), rtl)
    row = block(ws, 1, _("Group size"), ["", _("Vouchers")], pairs(c["sizes"]))
    block(ws, row, _("Nights by city"), ["", _("Nights")], pairs(c["cities"]))
    _autosize(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()

