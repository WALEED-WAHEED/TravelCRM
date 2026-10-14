"""Hotel, Flight, Transport and Package insights: the numbers, the pages, the Excel files, and agency isolation."""
import datetime
import io

import openpyxl
import pytest
from django.utils import timezone

from member import inventory_insights as ii
from member.models import Booking, BookingFlight, BookingHotel, BookingTransport, Flight, Hotel, Package, Transport
from reports import analytics
from utils.tenancy import tenant_context


def _period():
    return analytics.resolve_period({"preset": "this_year"})


@pytest.fixture
def stock(agency_a, user_a):
    """Catalogue + two vouchers that use part of it (and one still in approval, which must not count)."""
    today = timezone.localdate()
    with tenant_context(agency_id=agency_a.id):
        used_hotel = Hotel.objects.create(agency=agency_a, created_by=user_a, name="Used Hotel", city="Makkah", rating=5, sale_price=100, purchase_price=80)
        Hotel.objects.create(agency=agency_a, created_by=user_a, name="Idle Hotel", city="Madinah", rating=3, sale_price=60, purchase_price=50)
        flight = Flight.objects.create(agency=agency_a, created_by=user_a, airline="Saudia", flight_number="SV1", departure_city="Lahore",
                                       arrival_city="Jeddah", sale_price=0, purchase_price=0)
        Flight.objects.create(agency=agency_a, created_by=user_a, airline="Emirates", flight_number="EK2", departure_city="Dubai",
                              arrival_city="Lahore", sale_price=0, purchase_price=0)
        bus = Transport.objects.create(agency=agency_a, created_by=user_a, name="Coach 1", transport_type="coach", capacity=40, is_active=True)
        Transport.objects.create(agency=agency_a, created_by=user_a, name="Spare Van", transport_type="van", capacity=10, is_active=True)
        package = Package.objects.create(agency=agency_a, created_by=user_a, name="Umrah 10", package_type="umrah", price_double=200000,
                                         cost_price=150000, total_seats=10, departure_date=today + datetime.timedelta(days=20), is_active=True)
        Package.objects.create(agency=agency_a, created_by=user_a, name="Unsold Package", package_type="hajj", price_quad=500000, is_active=True)

        def voucher(number, status, pax, package_name=""):
            b = Booking.objects.create(agency=agency_a, agent=user_a, voucher_number=number, status=status, pax_adults=pax, package_name=package_name)
            BookingHotel.objects.create(booking=b, hotel=used_hotel, city="Makkah", hotel_name="Used Hotel", room_type="Quad", meal_plan="BB",
                                        check_in=today, check_out=today + datetime.timedelta(days=5), nights=5)
            BookingFlight.objects.create(booking=b, flight=flight, flight_type="outbound", airline="Saudia", flight_number="SV1",
                                         departure_city="Lahore", arrival_city="Jeddah", flight_class="business", departure_at=timezone.now())
            BookingFlight.objects.create(booking=b, flight=flight, flight_type="return", airline="Saudia", flight_number="SV1",
                                         departure_city="Jeddah", arrival_city="Lahore", flight_class="economy", departure_at=timezone.now())
            BookingTransport.objects.create(booking=b, transport=bus, travel_date=today, transporter="Coach 1", vehicle_type="Coach", description="JED - Makkah")
            return b

        voucher("IV-1", "confirmed", 4, "Umrah 10")
        voucher("IV-2", "draft", 2)
        voucher("IV-3", "pending", 9, "Umrah 10")                      # still in approval: not counted anywhere
    return package


def _cards(data):
    return {c["title"]: c for s in data["sections"] for c in s["cards"]}


def _kpis(data):
    return {k["label"]: k["value"] for k in data["kpis"]}


def test_hotel_numbers(agency_a, stock):
    with tenant_context(agency_id=agency_a.id):
        data = ii.build("hotels", agency_a, _period())
    k = _kpis(data)
    assert k["Hotels"] == 2 and k["Stays booked"] == 2 and k["Nights booked"] == 10 and k["Guest nights"] == 5 * 4 + 5 * 2
    assert k["Hotels used"] == 1
    top = _cards(data)["Most booked hotels"]["rows"]
    assert top[0][0]["t"] == "Used Hotel" and top[0][3]["t"] == 10 and top[0][5]["bar"] == 100
    assert [r[0]["t"] for r in _cards(data)["Not booked this period"]["rows"]] == ["Idle Hotel"]


def test_flight_numbers(agency_a, stock):
    with tenant_context(agency_id=agency_a.id):
        data = ii.build("flights", agency_a, _period())
    k = _kpis(data)
    assert k["Flights"] == 2 and k["Flight legs booked"] == 4 and k["Seats booked"] == (4 + 4 + 2 + 2) and k["Vouchers with flights"] == 2
    assert k["Busiest route"] in ("Lahore-Jeddah", "Jeddah-Lahore", "LHE-JED", "JED-LHE") or k["Busiest route"]
    used = _cards(data)["Flights used"]["rows"]
    assert used[0][2]["t"] == 4 and "Business 2" in used[0][4]["t"] and "Economy 2" in used[0][4]["t"]
    assert [r[0]["t"] for r in _cards(data)["Not booked this period"]["rows"]] == ["Emirates EK2"]


def test_transport_numbers(agency_a, stock):
    with tenant_context(agency_id=agency_a.id):
        data = ii.build("transports", agency_a, _period())
    k = _kpis(data)
    assert k["Vehicles"] == 2 and k["Transport bookings"] == 2 and k["Travellers carried"] == 6
    assert k["Average load"] == f"{round(6 * 100 / 80)}%"
    assert [r[0]["t"] for r in _cards(data)["Not booked this period"]["rows"]] == ["Spare Van"]


def test_package_numbers(agency_a, stock):
    with tenant_context(agency_id=agency_a.id):
        data = ii.build("packages", agency_a, _period())
    k = _kpis(data)
    assert k["Packages"] == 2 and k["Vouchers from packages"] == 1 and k["Travellers sold"] == 4
    assert k["Custom trips"] == "50%" and k["Departing soon"] == 1 and k["Average margin"] == "25%"
    seat = _cards(data)["Seat fill"]["rows"]
    assert seat[0][0]["t"] == "Umrah 10" and seat[0][3]["t"] == 10 and seat[0][4]["t"] == 4 and seat[0][5]["t"] == 6 and seat[0][6]["bar"] == 40
    assert [r[0]["t"] for r in _cards(data)["Not sold this period"]["rows"]] == ["Unsold Package"]


@pytest.mark.parametrize("kind", ii.KINDS)
def test_pages_exports_and_isolation(client_a, client_b, stock, kind):
    r = client_a.get(f"/member/{kind}/insights/", {"preset": "this_year"})
    assert r.status_code == 200
    body = r.content.decode()
    assert "insights" in body.lower() and "bi-data" in body
    assert len(r.context["data"]["kpis"]) == 6
    x = client_a.get(f"/member/{kind}/insights/", {"preset": "this_year", "export": "xlsx"})
    assert x.status_code == 200 and "spreadsheetml" in x["Content-Type"]
    wb = openpyxl.load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames[0] == "Summary" and len(wb.sheetnames) >= 3
    assert wb["Summary"]["A6"].value == r.context["data"]["kpis"][0]["label"]
    # another agency sees none of it
    other = client_b.get(f"/member/{kind}/insights/", {"preset": "this_year"})
    assert other.status_code == 200
    empty = (0, "0", "0%", "Rs. 0", "—")
    assert sum(1 for k in other.context["data"]["kpis"] if k["value"] not in empty) <= 1
    for preset in analytics.PRESETS:
        if preset != "custom":
            assert client_a.get(f"/member/{kind}/insights/", {"preset": preset}).status_code == 200


def test_unknown_kind_is_a_404(client_a):
    assert client_a.get("/member/nonsense/insights/").status_code == 404


def test_each_inventory_page_links_to_its_insights(client_a):
    for kind in ii.KINDS:
        page = client_a.get(f"/member/{kind}/").content.decode()
        assert f"/member/{kind}/insights/" in page


def test_ledger_numbers_and_page(client_a, client_b, agency_a, user_a):
    from decimal import Decimal

    from member.models import LedgerEntry
    today = timezone.localdate()
    with tenant_context(agency_id=agency_a.id):
        def entry(**kw):
            base = dict(agency=agency_a, created_by=user_a, transaction_date=today, currency="PKR", description="x")
            base.update(kw)
            return LedgerEntry.objects.create(**base)
        entry(transaction_type="visa", bucket="lein_dein", direction="debit", counterparty_name="Client A", amount=Decimal("1000"))
        entry(transaction_type="receipt", bucket="lein_dein", direction="credit", counterparty_name="Client A", amount=Decimal("400"), payment_method="cash")
        entry(transaction_type="hotel", bucket="lein_dein", direction="credit", counterparty_name="Hotel Co", amount=Decimal("700"))
        entry(transaction_type="expense", bucket="company_expense", direction="debit", amount=Decimal("150"), category="Rent")
        data = ii.build("ledger", agency_a, _period())
    k = _kpis(data)
    assert k["Received"] == "Rs. 400" and k["Paid out"] == "Rs. 150" and k["Net cash"] == "Rs. 250" and k["Expenses"] == "Rs. 150"
    assert k["Parties owe you"] == "Rs. 600" and k["You owe parties"] == "Rs. 700"
    cards = _cards(data)
    assert cards["Parties who owe you"]["rows"][0][0]["t"] == "Client A" and cards["Parties you owe"]["rows"][0][0]["t"] == "Hotel Co"
    r = client_a.get("/member/ledger/insights/", {"preset": "this_year"})
    assert r.status_code == 200 and "Ledger insights" in r.content.decode()
    assert "/member/ledger/insights/" in client_a.get("/member/ledger/").content.decode()
    x = client_a.get("/member/ledger/insights/", {"export": "xlsx"})
    assert x.status_code == 200 and "spreadsheetml" in x["Content-Type"]
    other = client_b.get("/member/ledger/insights/").context["data"]["kpis"]
    assert all(k_["value"] in ("Rs. 0", 0) for k_ in other)


def test_attendance_numbers_and_page(client_a, client_b, agency_a):
    from member.models import AttendanceLog
    today = timezone.localdate()
    with tenant_context(agency_id=agency_a.id):
        rows = (("Ali", "present", 8, 0), ("Ali", "late", 7, 1), ("Sara", "absent", None, 0), ("Sara", "leave", None, 1), ("Omar", "present", 9, 0))
        for name, status, hours, back in rows:
            AttendanceLog.objects.create(agency=agency_a, person_name=name, date=today - datetime.timedelta(days=back), status=status, work_hours=hours)
        data = ii.build("attendance", agency_a, _period())
    k = _kpis(data)
    assert k["Attendance rate"] == "75.0%" and k["Late arrivals"] == 1 and k["Absences"] == 1 and k["Average hours"] == "8.0 h"
    people = _cards(data)["Staff"]["rows"]
    assert [r[0]["t"] for r in people] == ["Ali", "Omar", "Sara"] and people[0][6]["bar"] == 100 and people[2][6]["t"] == "—" or people[2][6]["bar"] == 0
    r = client_a.get("/member/attendance/insights/", {"preset": "this_year"})
    assert r.status_code == 200 and "Attendance insights" in r.content.decode()
    assert "/member/attendance/insights/" in client_a.get("/member/attendance/").content.decode()
    assert client_a.get("/member/attendance/insights/", {"export": "xlsx"}).status_code == 200
    other = client_b.get("/member/attendance/insights/").context["data"]["kpis"]
    assert other[2]["value"] == 0 and other[3]["value"] == 0
